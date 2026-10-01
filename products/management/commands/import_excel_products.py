from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from dealers.models import Group
from products.models import Product


class Command(BaseCommand):
    help = "Import or update the product master from the New zealand dairy workbook sheet."

    def add_arguments(self, parser):
        parser.add_argument("--file", default=str(settings.BASE_DIR / "reference" / "Daily Sale Calculate All GR.xlsm"))
        parser.add_argument("--sheet", default="New zealand dairy")
        parser.add_argument("--dry-run", action="store_true")

    @staticmethod
    def clean_text(value):
        return str(value).strip() if value not in (None, "") else ""

    CATEGORY_MAP = {"Mlk": "Milk", "wefr": "Wefr", "OIL": "Oil"}

    @staticmethod
    def money(value):
        if value in (None, "", " "):
            return Decimal("0")
        try:
            return Decimal(str(value).strip())
        except (InvalidOperation, ValueError) as exc:
            raise CommandError(f"Invalid numeric product value: {value!r}") from exc

    def handle(self, *args, **options):
        from openpyxl import load_workbook

        workbook_path = Path(options["file"])
        if not workbook_path.exists():
            raise CommandError(f"Workbook not found: {workbook_path}")
        workbook = load_workbook(workbook_path, data_only=True, read_only=True, keep_vba=True)
        if options["sheet"] not in workbook.sheetnames:
            raise CommandError(f"Sheet not found: {options['sheet']}")
        sheet = workbook[options["sheet"]]
        group = Group.objects.get(name="New Zealand Dairy")
        imported = updated = skipped = 0
        with transaction.atomic():
            current_category = ""
            for row in range(4, 85):
                name = self.clean_text(sheet.cell(row, 3).value)
                if not name:
                    skipped += 1
                    continue
                raw_category = self.clean_text(sheet.cell(row, 2).value)
                if raw_category:
                    current_category = raw_category
                category = self.CATEGORY_MAP.get(current_category, current_category)
                values = {
                    "category": category,
                    "ctn_size": max(int(self.money(sheet.cell(row, 4).value)), 1),
                    "dealer_price": self.money(sheet.cell(row, 5).value),
                    "trade_price": self.money(sheet.cell(row, 6).value),
                    "foc": self.money(sheet.cell(row, 7).value),
                    "selling_price": self.money(sheet.cell(row, 8).value),
                    "reference_dealer_price": self.money(sheet.cell(row, 21).value),
                    "markup_price": self.money(sheet.cell(row, 22).value),
                }
                product = Product.objects.filter(group=group, name__iexact=name, category__name__iexact=category).first()
                if product:
                    for field, value in values.items():
                        setattr(product, field, value)
                    if not options["dry_run"]:
                        product.save(update_fields=list(values))
                    updated += 1
                else:
                    if not options["dry_run"]:
                        Product.objects.create(group=group, name=name, **values)
                    imported += 1
            if options["dry_run"]:
                transaction.set_rollback(True)
        self.stdout.write(self.style.SUCCESS(f"Imported: {imported}; updated: {updated}; skipped: {skipped}"))
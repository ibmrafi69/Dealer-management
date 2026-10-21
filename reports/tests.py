from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.test import TestCase
from openpyxl import load_workbook

from dealers.models import Dealer, Group, Route
from products.models import Product
from sales.models import PaymentReconciliation, SalesItem, SalesTransaction

from . import views


class ReportTests(TestCase):
    def test_monthly_export_returns_xlsx(self):
        dealer = Dealer.objects.create(name="Report Dealer", code="REPORT")
        product = Product.objects.create(name="Report Product", trade_price=10)
        sale = SalesTransaction.objects.create(dealer=dealer, date=date(2026, 9, 2), invoice_number="R-1")
        SalesItem.objects.create(transaction=sale, product=product, order_quantity=2)
        response = self.client.get("/reports/monthly/export/?year=2026&month=9")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_daily_export_returns_xlsx(self):
        response = self.client.get("/reports/daily/export/?date=2026-09-02")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_monthly_report_keeps_saved_shortage_with_sales_return_and_damage(self):
        dealer = Dealer.objects.create(name="Monthly Dealer", code="MONTHLY")
        product = Product.objects.create(name="Monthly Product", trade_price=100, selling_price=100)
        sale = SalesTransaction.objects.create(dealer=dealer, date=date(2026, 8, 15), invoice_number="M-1")
        SalesItem.objects.create(
            transaction=sale,
            product=product,
            order_quantity=10,
            return_quantity=2,
            damage_quantity=1,
        )
        reconciliation = PaymentReconciliation.objects.create(transaction=sale, cash=500, credit=100)

        response = self.client.get("/reports/monthly/?year=2026&month=8")

        self.assertEqual(response.status_code, 200)
        reconciliation.refresh_from_db()
        self.assertEqual(reconciliation.shortage, Decimal("100.00"))
        self.assertEqual(response.context["totals"]["sales"], Decimal("700.00"))
        self.assertEqual(response.context["totals"]["returns"], Decimal("200.00"))
        self.assertEqual(response.context["totals"]["damage"], Decimal("100.00"))
        self.assertEqual(response.context["totals"]["shortage"], Decimal("100.00"))
        self.assertContains(response, "Shortage")

    def test_monthly_summary_and_exports_are_finalized_only_and_route_filtered(self):
        selected_route = Route.objects.create(route_name="Export Route A", weekday=Route.MONDAY)
        other_route = Route.objects.create(route_name="Export Route B", weekday=Route.TUESDAY)
        product = Product.objects.create(name="Export Product", ctn_size=12, trade_price=100, selling_price=100)
        finalized = SalesTransaction.objects.create(
            route=selected_route,
            date=date(2026, 10, 12),
            invoice_number="EXPORT-FINAL",
            status=SalesTransaction.Status.FINALIZED,
        )
        SalesItem.objects.create(transaction=finalized, product=product, order_quantity=10, return_quantity=2, damage_quantity=1)
        PaymentReconciliation.objects.create(transaction=finalized, cash=500)
        pending = SalesTransaction.objects.create(
            route=selected_route,
            date=date(2026, 10, 12),
            invoice_number="EXPORT-PENDING",
            status=SalesTransaction.Status.PENDING,
        )
        SalesItem.objects.create(transaction=pending, product=product, order_quantity=8)
        other_route_sale = SalesTransaction.objects.create(
            route=other_route,
            date=date(2026, 10, 13),
            invoice_number="EXPORT-OTHER-ROUTE",
            status=SalesTransaction.Status.FINALIZED,
        )
        SalesItem.objects.create(transaction=other_route_sale, product=product, order_quantity=4)

        filters = f"?year=2026&month=10&route={selected_route.pk}"
        summary = self.client.get(f"/daily-summary/{filters}")
        self.assertEqual(summary.status_code, 200)
        self.assertContains(summary, "EXPORT-FINAL")
        self.assertNotContains(summary, "EXPORT-PENDING")
        self.assertNotContains(summary, "EXPORT-OTHER-ROUTE")
        self.assertContains(summary, f'href="/sales/transactions/{finalized.pk}/"')
        self.assertNotContains(summary, f'/sales/transactions/{finalized.pk}/edit/')
        # Daily Sales quantities are stored as individual pieces. CTN size is
        # retained for stock receipts and invoice history, not conversion here.
        self.assertContains(summary, "10.00 / 10.00")

        invoice = self.client.get(f"/sales/transactions/{finalized.pk}/")
        self.assertEqual(invoice.status_code, 200)
        self.assertContains(invoice, "Read-only invoice detail")
        # The shared sidebar has a group-switcher form; a finalized invoice
        # must only omit its own editing endpoint.
        self.assertNotContains(invoice, f'action="/sales/transactions/{finalized.pk}/edit/"')
        self.assertEqual(self.client.get(f"/sales/transactions/{pending.pk}/").status_code, 404)

        excel = self.client.get(f"/reports/monthly/export/{filters}")
        self.assertEqual(excel.status_code, 200)
        self.assertEqual(excel["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        workbook = load_workbook(BytesIO(excel.content), data_only=True)
        invoice_sheet = workbook["Invoice Records"]
        exported_invoices = [invoice_sheet.cell(row=row, column=2).value for row in range(5, invoice_sheet.max_row + 1)]
        self.assertIn("EXPORT-FINAL", exported_invoices)
        self.assertNotIn("EXPORT-PENDING", exported_invoices)
        self.assertNotIn("EXPORT-OTHER-ROUTE", exported_invoices)
        self.assertEqual(invoice_sheet.cell(row=5, column=6).value, 10)

        pdf = self.client.get(f"/reports/monthly/export/pdf/{filters}")
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF-"))

        invoices_pdf = self.client.get(f"/reports/monthly/export/invoices/pdf/{filters}")
        self.assertEqual(invoices_pdf.status_code, 200)
        self.assertEqual(invoices_pdf["Content-Type"], "application/pdf")
        self.assertEqual(invoices_pdf["Content-Disposition"], 'attachment; filename="aafia-invoices-2026-10.pdf"')
        self.assertTrue(invoices_pdf.content.startswith(b"%PDF-"))

    def test_monthly_invoices_pdf_export_uses_the_shared_finalized_group_filter(self):
        group = Group.objects.get(name="New Zealand Dairy")
        route = Route.objects.create(group=group, route_name="Invoice PDF Route", weekday=Route.MONDAY)
        product = Product.objects.create(group=group, name="Invoice PDF Product", trade_price=100, selling_price=100)
        sale = SalesTransaction.objects.create(
            group=group,
            route=route,
            date=date(2026, 10, 12),
            invoice_number="INVOICES-PDF-FINAL",
            status=SalesTransaction.Status.FINALIZED,
        )
        SalesItem.objects.create(group=group, transaction=sale, product=product, order_quantity=2)

        summary = self.client.get(f"/daily-summary/?year=2026&month=10&route={route.pk}")
        self.assertContains(summary, f"/reports/monthly/export/invoices/pdf/?year=2026&amp;month=10&amp;route={route.pk}")

        with patch("reports.views._filtered_transactions", wraps=views._filtered_transactions) as filtered:
            response = self.client.get(f"/reports/monthly/export/invoices/pdf/?year=2026&month=10&route={route.pk}")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertEqual(response["Content-Disposition"], 'attachment; filename="aafia-invoices-2026-10.pdf"')
        self.assertTrue(response.content.startswith(b"%PDF-"))
        filtered.assert_called_once()

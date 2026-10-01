from calendar import month_name
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.shortcuts import render
from django.utils.html import escape

from dealers.context import get_current_group
from dealers.models import Route
from sales.models import SalesTransaction


def _filtered_transactions(request):
    group = get_current_group(request)
    today = date.today()
    try:
        year = int(request.GET.get("year", today.year))
        month = int(request.GET.get("month", today.month))
        if not 1 <= month <= 12:
            raise ValueError
    except ValueError:
        year, month = today.year, today.month
    route_id = request.GET.get("route", "")
    transactions = (
        SalesTransaction.objects.filter(
            group=group,
            status=SalesTransaction.Status.FINALIZED,
            date__year=year,
            date__month=month,
        )
        .select_related("dealer", "route", "reconciliation")
        .prefetch_related("items__product", "damages")
    )
    if route_id:
        transactions = transactions.filter(route_id=route_id)
    return group, year, month, route_id, transactions


def _invoice_records(transactions):
    """Expose saved finalized-invoice totals shared by both export formats."""
    return [
        {
            "date": sale.date,
            "invoice": sale.invoice_number,
            "route": sale.route_display,
            "day": sale.route_day or sale.date.strftime("%A"),
            "status": sale.get_status_display(),
            "sent_quantity": sale.sent_quantity_total,
            "sent_pieces": sale.sent_pieces_total,
            "sale_quantity": sale.sale_quantity_total,
            "sale_pieces": sale.sale_pieces_total,
            "return_quantity": sale.return_quantity_total,
            "return_pieces": sale.return_pieces_total,
            "damage_quantity": sale.damage_quantity_total,
            "damage_pieces": sale.damage_pieces_total,
            "sales": sale.sale_total,
            "returns": sale.return_total,
            "damage": sale.damage_total,
            "discount": sale.total_discount,
            "expenses": getattr(getattr(sale, "reconciliation", None), "expenses", Decimal("0")),
            "shortage": getattr(getattr(sale, "reconciliation", None), "shortage", sale.grand_total_sale),
            "net": sale.net_sale,
        }
        for sale in transactions
    ]


def _invoice_totals(records):
    fields = ("sales", "returns", "damage", "discount", "expenses", "shortage", "net")
    return {field: sum((record[field] for record in records), Decimal("0")) for field in fields}


def _route_label(route_id, group):
    if not route_id:
        return "All routes"
    route = Route.objects.filter(pk=route_id, group=group).only("route_name").first()
    return route.route_name if route else "Selected route"


def monthly_report(request):
    group, year, month, route_id, transactions = _filtered_transactions(request)
    daily = {}
    product_totals = {}
    for sale in transactions:
        summary = daily.setdefault(
            sale.date,
            {
                "sales": Decimal("0"),
                "returns": Decimal("0"),
                "damage": Decimal("0"),
                "discount": Decimal("0"),
                "expenses": Decimal("0"),
                "shortage": Decimal("0"),
                "net": Decimal("0"),
            },
        )
        summary["sales"] += sale.sale_total
        summary["returns"] += sale.return_total
        summary["damage"] += sale.damage_total
        summary["discount"] += sale.total_discount + sale.free_product_total
        summary["expenses"] += getattr(getattr(sale, "reconciliation", None), "expenses", Decimal("0"))
        summary["shortage"] += getattr(getattr(sale, "reconciliation", None), "shortage", sale.grand_total_sale)
        summary["net"] += sale.net_sale
        for item in sale.items.all():
            key = (
                item.product_id_snapshot if item.product_id_snapshot is not None else item.product_id,
                item.product_name_snapshot or (item.product.name if item.product_id else ""),
                item.category_snapshot,
            )
            row = product_totals.setdefault(
                key,
                {
                    "name": key[1],
                    "category": key[2],
                    "order": Decimal("0"),
                    "return": Decimal("0"),
                    "sale": Decimal("0"),
                    "damage": Decimal("0"),
                    "value": Decimal("0"),
                },
            )
            row["order"] += item.order_quantity
            row["return"] += item.return_quantity
            row["sale"] += item.sale_quantity
            row["damage"] += item.damage_quantity
            row["value"] += item.sale_value
    daily_records = _invoice_records(transactions)
    totals = _invoice_totals(daily_records)
    return render(
        request,
        "reports/monthly.html",
        {
            "daily": sorted(daily.items()),
            "daily_records": daily_records,
            "product_totals": product_totals.items(),
            "totals": totals,
            "routes": Route.objects.filter(group=group, active=True),
            "year": year,
            "month": month,
            "month_name": month_name[month],
            "month_options": [(number, month_name[number]) for number in range(1, 13)],
            "route_id": route_id,
        },
    )


def _style_header(sheet, row, header_fill, header_font):
    for cell in sheet[row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = cell.alignment.copy(horizontal="center", vertical="center", wrap_text=True)


def _fit_columns(sheet, maximum=28):
    from openpyxl.utils import get_column_letter

    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, maximum)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width


def export_monthly(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    group, year, month, route_id, transactions = _filtered_transactions(request)
    records = _invoice_records(transactions)
    totals = _invoice_totals(records)
    route_label = _route_label(route_id, group)

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Invoice Records"
    details_sheet = workbook.create_sheet("Product Details")
    products_sheet = workbook.create_sheet("Product Summary")
    header_fill = PatternFill("solid", fgColor="123A5A")
    title_fill = PatternFill("solid", fgColor="E8F1F7")
    header_font = Font(color="FFFFFF", bold=True)

    headers = [
        "Date",
        "Invoice Number",
        "Route",
        "Day",
        "Sent Qty",
        "Sent Pieces",
        "Sale Qty",
        "Sale Pieces",
        "Return Qty",
        "Return Pieces",
        "Damage Qty",
        "Damage Pieces",
        "Total Sale Amount",
        "Total Discount",
        "Total Shortage",
        "Net Sale",
        "Status",
    ]
    summary_sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    summary_sheet["A1"] = "AAFIA ENTERPRISE - MONTHLY SALES SUMMARY"
    summary_sheet["A1"].font = Font(size=15, bold=True, color="123A5A")
    summary_sheet["A1"].fill = title_fill
    summary_sheet["A1"].alignment = Alignment(horizontal="center")
    summary_sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    summary_sheet["A2"] = f"Selected: {month_name[month]} {year} | Route: {route_label} | Finalized invoices only"
    summary_sheet["A2"].alignment = Alignment(horizontal="center")
    summary_sheet.append([])
    summary_sheet.append(headers)
    _style_header(summary_sheet, 4, header_fill, header_font)

    for record in records:
        summary_sheet.append(
            [
                record["date"],
                record["invoice"],
                record["route"],
                record["day"],
                record["sent_quantity"],
                record["sent_pieces"],
                record["sale_quantity"],
                record["sale_pieces"],
                record["return_quantity"],
                record["return_pieces"],
                record["damage_quantity"],
                record["damage_pieces"],
                record["sales"],
                record["discount"],
                record["shortage"],
                record["net"],
                record["status"],
            ]
        )
    summary_sheet.freeze_panes = "A5"
    for row in summary_sheet.iter_rows(min_row=5, max_row=summary_sheet.max_row):
        row[0].number_format = "dd mmm yyyy"
        for cell in row[4:16]:
            cell.number_format = "#,##0.00"

    totals_title_row = summary_sheet.max_row + 2
    summary_sheet.merge_cells(start_row=totals_title_row, start_column=1, end_row=totals_title_row, end_column=2)
    summary_sheet.cell(totals_title_row, 1, "MONTHLY TOTALS")
    summary_sheet.cell(totals_title_row, 1).fill = title_fill
    summary_sheet.cell(totals_title_row, 1).font = Font(bold=True, color="123A5A")
    for label, field in (
        ("Total Sales", "sales"),
        ("Total Returns", "returns"),
        ("Total Damage", "damage"),
        ("Total Discount", "discount"),
        ("Total Shortage", "shortage"),
        ("Total Net Sales", "net"),
    ):
        summary_sheet.append([label, totals[field]])
        summary_sheet.cell(summary_sheet.max_row, 2).number_format = "#,##0.00"

    details_headers = [
        "Date", "Route", "Invoice", "Day", "Category", "Product", "CTN Size",
        "Dealer Price", "Trade Price", "FOC Price", "Sale Price", "Order Qty",
        "Return Qty", "Sale Qty", "Damage Qty", "Free Product", "Line Discount",
        "Order Value", "Return Value", "Damage Value", "Free Value", "Sale Value",
    ]
    details_sheet.append(details_headers)
    _style_header(details_sheet, 1, header_fill, header_font)
    for sale in transactions:
        for item in sale.items.all():
            details_sheet.append(
                [
                    sale.date, sale.route_display, sale.invoice_number, sale.route_day or sale.date.strftime("%A"),
                    item.category_snapshot, item.product_name_snapshot, item.ctn_size_snapshot,
                    item.dealer_price_snapshot, item.trade_price_snapshot, item.foc_snapshot,
                    item.sale_price_snapshot, item.order_quantity, item.return_quantity,
                    item.sale_quantity, item.damage_quantity, item.free_product_quantity,
                    item.discount, item.order_value, item.return_value, item.damage_value,
                    item.free_product_value, item.sale_value,
                ]
            )
    details_sheet.freeze_panes = "A2"
    for row in details_sheet.iter_rows(min_row=2, max_row=details_sheet.max_row):
        row[0].number_format = "dd mmm yyyy"
        for cell in row[6:]:
            cell.number_format = "#,##0.00"

    products_sheet.append(["Category", "Product", "Total Sent Qty", "Total Return Qty", "Total Sale Qty", "Total Damage Qty", "Total Sale Value"])
    _style_header(products_sheet, 1, header_fill, header_font)
    product_rows = {}
    for sale in transactions:
        for item in sale.items.all():
            key = (
                item.product_id_snapshot if item.product_id_snapshot is not None else item.product_id,
                item.category_snapshot,
                item.product_name_snapshot or (item.product.name if item.product_id else ""),
            )
            row = product_rows.setdefault(key, [Decimal("0")] * 5)
            row[0] += item.order_quantity
            row[1] += item.return_quantity
            row[2] += item.sale_quantity
            row[3] += item.damage_quantity
            row[4] += item.sale_value
    for (_, category, product), values in sorted(product_rows.items()):
        products_sheet.append([category, product, *values])
    products_sheet.freeze_panes = "A2"
    for row in products_sheet.iter_rows(min_row=2, max_row=products_sheet.max_row):
        for cell in row[2:]:
            cell.number_format = "#,##0.00"

    _fit_columns(summary_sheet, maximum=24)
    _fit_columns(details_sheet, maximum=28)
    _fit_columns(products_sheet, maximum=28)
    summary_sheet.column_dimensions["A"].width = 14
    summary_sheet.column_dimensions["B"].width = 22
    summary_sheet.column_dimensions["C"].width = 18

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="aafia-monthly-{year}-{month:02d}.xlsx"'
    return response


def _format_money(value):
    return f"{value:,.2f}"


def _format_quantity_pair(record, prefix):
    return f"{record[f'{prefix}_quantity']:.2f} / {record[f'{prefix}_pieces']:.2f}"


def export_monthly_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    group, year, month, route_id, transactions = _filtered_transactions(request)
    records = _invoice_records(transactions)
    totals = _invoice_totals(records)
    route_label = _route_label(route_id, group)
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title=f"AAFIA Enterprise Monthly Sales Summary - {month_name[month]} {year}",
    )
    styles = getSampleStyleSheet()
    brand_style = ParagraphStyle(
        "Brand",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        textColor=colors.HexColor("#123A5A"),
        alignment=TA_CENTER,
        spaceAfter=3,
    )
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        alignment=TA_CENTER,
        spaceAfter=6,
    )
    detail_style = ParagraphStyle(
        "ReportDetail",
        parent=styles["Normal"],
        fontSize=8,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#394B59"),
        spaceAfter=10,
    )
    story = [
        Paragraph("AAFIA ENTERPRISE", brand_style),
        Paragraph("MONTHLY SALES SUMMARY", title_style),
        Paragraph(
            f"Selected: {month_name[month]} {year} | Route: {escape(route_label)} | Finalized invoices only",
            detail_style,
        ),
    ]
    report_rows = [["Date", "Invoice", "Route", "Sent\nQty / Pcs", "Sale\nQty / Pcs", "Return\nQty / Pcs", "Damage\nQty / Pcs", "Total Tk"]]
    for record in records:
        report_rows.append(
            [
                record["date"].strftime("%d %b %Y"),
                record["invoice"],
                record["route"],
                _format_quantity_pair(record, "sent"),
                _format_quantity_pair(record, "sale"),
                _format_quantity_pair(record, "return"),
                _format_quantity_pair(record, "damage"),
                _format_money(record["sales"]),
            ]
        )
    if len(report_rows) == 1:
        report_rows.append(["No finalized invoices for the selected filters.", "", "", "", "", "", "", ""])
    report_table = Table(report_rows, colWidths=[23 * mm, 31 * mm, 32 * mm, 25 * mm, 25 * mm, 25 * mm, 25 * mm, 28 * mm], repeatRows=1)
    report_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123A5A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B8C7D1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F8FB")]),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    if not records:
        report_table.setStyle(TableStyle([("SPAN", (0, 1), (-1, 1)), ("ALIGN", (0, 1), (-1, 1), "LEFT")]))
    story.extend([report_table, Spacer(1, 10)])

    totals_rows = [["MONTHLY TOTALS", "Amount (Tk)"]]
    totals_rows.extend(
        [
            ["Total Sales", _format_money(totals["sales"])],
            ["Total Returns", _format_money(totals["returns"])],
            ["Total Damage", _format_money(totals["damage"])],
            ["Total Discount", _format_money(totals["discount"])],
            ["Total Shortage", _format_money(totals["shortage"])],
            ["Total Net Sales", _format_money(totals["net"])],
        ]
    )
    totals_table = Table(totals_rows, colWidths=[55 * mm, 35 * mm], hAlign="RIGHT")
    totals_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCEAF3")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#123A5A")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B8C7D1")),
                ("ALIGN", (1, 1), (1, -1), "RIGHT"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(totals_table)

    def draw_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#627D8A"))
        canvas.drawString(doc.leftMargin, 8 * mm, "AAFIA ENTERPRISE - Monthly Sales Summary")
        canvas.drawRightString(landscape(A4)[0] - doc.rightMargin, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="aafia-monthly-{year}-{month:02d}.pdf"'
    return response


def export_daily(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    try:
        selected_date = date.fromisoformat(request.GET.get("date", ""))
    except ValueError:
        selected_date = date.today()
    group = get_current_group(request)
    transactions = SalesTransaction.objects.filter(group=group, status=SalesTransaction.Status.FINALIZED, date=selected_date).select_related("dealer", "route").prefetch_related("items__product", "damages")
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Daily Summary"
    details = workbook.create_sheet("Daily Details")
    reconciliation = workbook.create_sheet("Reconciliation")
    header_fill = PatternFill("solid", fgColor="123A5A")
    header_font = Font(color="FFFFFF", bold=True)

    def style_header(sheet):
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"

    summary.append(["Date", "Total Sale", "Damage", "Grand Total Sale", "Discount", "Free Product Value", "Net Sale"])
    for sale in transactions:
        summary.append([selected_date, sale.sale_total, sale.damage_total, sale.grand_total_sale, sale.total_discount, sale.free_product_total, sale.net_sale])
    style_header(summary)
    details.append(["Date", "Route", "Invoice", "Category", "Product", "CTN Size", "Dealer Price", "Trade Price", "FOC", "Sale Price", "Order Qty", "Return Qty", "Sale Qty", "Value Price", "Damage Qty", "Damage Value", "Free Product", "Free P. Value", "Order Value", "Return Value", "Sale Value"])
    for sale in transactions:
        for item in sale.items.all():
            details.append([selected_date, sale.route_display, sale.invoice_number, item.category_snapshot, item.product_name_snapshot, item.ctn_size_snapshot, item.dealer_price_snapshot, item.trade_price_snapshot, item.foc_quantity, item.sale_price_snapshot, item.order_quantity, item.return_quantity, item.sale_quantity, item.unit_price, item.damage_quantity, item.damage_value, item.free_product_quantity, item.free_product_value, item.order_value, item.return_value, item.sale_value])
    style_header(details)
    reconciliation.append(["Invoice", "Grand Total Sale", "Cash", "Expenses", "Credit", "Discount", "Others", "Shortage / Variance"])
    for sale in transactions:
        payment = getattr(sale, "reconciliation", None)
        reconciliation.append([sale.invoice_number, sale.grand_total_sale, payment.cash if payment else 0, payment.expenses if payment else 0, payment.credit if payment else 0, payment.discount if payment else 0, payment.others if payment else 0, payment.shortage if payment else sale.grand_total_sale])
    style_header(reconciliation)
    for sheet in workbook.worksheets:
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 28)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (Decimal, float, int)) and cell.column > 1:
                    cell.number_format = "#,##0.00"
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="aafia-daily-{selected_date.isoformat()}.xlsx"'
    return response
